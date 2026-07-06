# tests/test_sheet_management.py
"""create_sheet(index=...), move_sheet, copy_worksheet, and move_range."""

import io

import openpyxl
import pytest

from openpyxl_rust import Workbook
from openpyxl_rust.styles import Font
from openpyxl_rust.worksheet import _translate_formula


def _readback(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


class TestCreateSheetIndex:
    def test_insert_first(self):
        wb = Workbook()
        wb.active.title = "B"
        wb.create_sheet("A", index=0)
        assert wb.sheetnames == ["A", "B"]
        assert _readback(wb).sheetnames == ["A", "B"]

    def test_insert_middle(self):
        wb = Workbook()
        wb.active.title = "A"
        wb.create_sheet("C")
        wb.create_sheet("B", index=1)
        assert _readback(wb).sheetnames == ["A", "B", "C"]

    def test_values_follow_sheets_after_insert(self):
        wb = Workbook()
        wb.active["A1"] = "original"
        inserted = wb.create_sheet("Inserted", index=0)
        inserted["A1"] = "new"
        rb = _readback(wb)
        assert rb["Sheet"]["A1"].value == "original"
        assert rb["Inserted"]["A1"].value == "new"


class TestMoveSheet:
    def test_move_by_name(self):
        wb = Workbook()
        wb.active.title = "A"
        wb.create_sheet("B")
        wb.create_sheet("C")
        wb.move_sheet("C", offset=-2)
        assert wb.sheetnames == ["C", "A", "B"]
        assert _readback(wb).sheetnames == ["C", "A", "B"]

    def test_move_by_object_forward(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "A"
        wb.create_sheet("B")
        wb.move_sheet(ws, offset=1)
        assert _readback(wb).sheetnames == ["B", "A"]

    def test_move_unknown_sheet_raises(self):
        wb = Workbook()
        with pytest.raises(KeyError):
            wb.move_sheet("nope", offset=1)

    def test_values_follow_after_move(self):
        wb = Workbook()
        wb.active.title = "A"
        wb.active["A1"] = "in A"
        b = wb.create_sheet("B")
        b["A1"] = "in B"
        wb.move_sheet("B", offset=-1)
        rb = _readback(wb)
        assert rb["A"]["A1"].value == "in A"
        assert rb["B"]["A1"].value == "in B"


class TestCopyWorksheet:
    def test_copy_values_and_formats(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "x"
        ws["A1"].font = Font(bold=True)
        ws["B2"] = 3.5
        copy = wb.copy_worksheet(ws)
        assert copy.title == "Sheet Copy"
        rb = _readback(wb)
        assert rb["Sheet Copy"]["A1"].value == "x"
        assert rb["Sheet Copy"]["A1"].font.bold is True
        assert rb["Sheet Copy"]["B2"].value == 3.5

    def test_copy_is_independent(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "original"
        copy = wb.copy_worksheet(ws)
        copy["A1"] = "changed"
        rb = _readback(wb)
        assert rb["Sheet"]["A1"].value == "original"
        assert rb["Sheet Copy"]["A1"].value == "changed"

    def test_copy_of_copy_gets_unique_title(self):
        wb = Workbook()
        c1 = wb.copy_worksheet(wb.active)
        c2 = wb.copy_worksheet(wb.active)
        assert c1.title != c2.title
        assert len(set(wb.sheetnames)) == 3

    def test_copy_dimensions_and_merges(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "m"
        ws.merge_cells("A1:B2")
        ws.column_dimensions["A"].width = 33
        ws.row_dimensions[1].height = 40
        copy = wb.copy_worksheet(ws)
        assert copy.column_dimensions["A"].width == 33
        assert copy.row_dimensions[1].height == 40
        rb = _readback(wb)
        assert any(str(m) == "A1:B2" for m in rb["Sheet Copy"].merged_cells.ranges)

    def test_copy_foreign_sheet_raises(self):
        wb1 = Workbook()
        wb2 = Workbook()
        with pytest.raises(ValueError):
            wb1.copy_worksheet(wb2.active)


class TestMoveRange:
    def test_basic_move(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws["B2"] = 2
        ws.move_range("A1:B2", rows=1, cols=1)
        assert ws["A1"].value is None
        assert ws["B2"].value == 1
        assert ws["C3"].value == 2

    def test_move_formats_travel(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "styled"
        ws["A1"].font = Font(italic=True)
        ws.move_range("A1:A1", rows=0, cols=3)
        assert ws["D1"].font.italic is True
        rb = _readback(wb)
        assert rb.active["D1"].font.italic is True
        assert rb.active["A1"].value is None

    def test_destination_overwritten(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "mover"
        ws["B1"] = "victim"
        ws["B1"].font = Font(bold=True)
        ws.move_range("A1:A1", rows=0, cols=1)
        assert ws["B1"].value == "mover"
        assert ws["B1"].font is None

    def test_move_off_sheet_raises(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        with pytest.raises(ValueError):
            ws.move_range("A1:A1", rows=-1, cols=0)

    def test_noop_move(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws.move_range("A1:A1", rows=0, cols=0)
        assert ws["A1"].value == 1


class TestFormulaTranslation:
    @pytest.mark.parametrize(
        ("formula", "rows", "cols", "expected"),
        [
            ("=A1", 1, 0, "=A2"),
            ("=A1", 0, 1, "=B1"),
            ("=$A$1", 5, 5, "=$A$1"),
            ("=$A1", 1, 1, "=$A2"),
            ("=A$1", 1, 1, "=B$1"),
            ("=SUM(A1:B2)", 1, 1, "=SUM(B2:C3)"),
            ("=A1+B2*C3", 1, 0, "=A2+B3*C4"),
            ('=IF(A1>0,"yes A1","no")', 0, 1, '=IF(B1>0,"yes A1","no")'),
            ("=LOG10(A1)", 1, 0, "=LOG10(A2)"),
            ("=Sheet2!A1", 1, 1, "=Sheet2!B2"),
            ("=A1", -1, 0, "=#REF!"),
        ],
    )
    def test_translate(self, formula, rows, cols, expected):
        assert _translate_formula(formula, rows, cols) == expected
