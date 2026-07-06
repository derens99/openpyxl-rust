# tests/test_named_styles.py
"""NamedStyle: registration, application by name/object, openpyxl read-back."""

import openpyxl as real_openpyxl
import pytest

from openpyxl_rust import Workbook
from openpyxl_rust.styles import (
    Alignment,
    Border,
    Font,
    NamedStyle,
    PatternFill,
    Protection,
    Side,
)


class TestRegistration:
    def test_default_named_styles(self):
        wb = Workbook()
        assert wb.named_styles == ["Normal"]

    def test_add_named_style(self):
        wb = Workbook()
        wb.add_named_style(NamedStyle(name="highlight", font=Font(bold=True)))
        assert wb.named_styles == ["Normal", "highlight"]

    def test_duplicate_name_raises(self):
        wb = Workbook()
        wb.add_named_style(NamedStyle(name="highlight"))
        with pytest.raises(ValueError, match="exists already"):
            wb.add_named_style(NamedStyle(name="highlight"))

    def test_assigning_style_object_auto_registers(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws["A1"].style = NamedStyle(name="auto", font=Font(italic=True))
        assert "auto" in wb.named_styles
        assert ws["A1"].style == "auto"

    def test_unknown_style_name_raises(self):
        wb = Workbook()
        ws = wb.active
        with pytest.raises(ValueError, match="is not a known style"):
            ws["A1"].style = "nope"


class TestApplication:
    def test_default_cell_style_is_normal(self):
        wb = Workbook()
        assert wb.active["A1"].style == "Normal"

    def test_style_getter_returns_name(self):
        wb = Workbook()
        ws = wb.active
        wb.add_named_style(NamedStyle(name="highlight"))
        ws["A1"].style = "highlight"
        assert ws["A1"].style == "highlight"

    def test_components_applied_to_cell(self):
        wb = Workbook()
        ws = wb.active
        style = NamedStyle(
            name="full",
            font=Font(bold=True, size=14),
            fill=PatternFill(fill_type="solid", start_color="FFFF00"),
            border=Border(bottom=Side(style="thin", color="000000")),
            alignment=Alignment(horizontal="center"),
            number_format="0.00%",
            protection=Protection(locked=False),
        )
        wb.add_named_style(style)
        ws["A1"] = 0.5
        ws["A1"].style = "full"
        c = ws["A1"]
        assert c.font.bold is True
        assert c.font.size == 14
        assert c.fill.start_color == "FFFF00"
        assert c.border.bottom.style == "thin"
        assert c.alignment.horizontal == "center"
        assert c.number_format == "0.00%"
        assert c.protection.locked is False

    def test_partial_style_leaves_other_components(self):
        wb = Workbook()
        ws = wb.active
        wb.add_named_style(NamedStyle(name="bold_only", font=Font(bold=True)))
        ws["A1"] = 1
        ws["A1"].number_format = "0.00"
        ws["A1"].style = "bold_only"
        assert ws["A1"].font.bold is True
        assert ws["A1"].number_format == "0.00"

    def test_direct_override_after_style_wins(self):
        wb = Workbook()
        ws = wb.active
        wb.add_named_style(NamedStyle(name="highlight", font=Font(bold=True, size=14)))
        ws["A1"].style = "highlight"
        ws["A1"].font = Font(bold=False, size=9)
        assert ws["A1"].font.size == 9
        assert ws["A1"].style == "highlight"

    def test_style_persists_across_cell_lookups(self):
        wb = Workbook()
        ws = wb.active
        wb.add_named_style(NamedStyle(name="highlight", font=Font(bold=True)))
        ws.cell(row=2, column=3).style = "highlight"
        assert ws.cell(row=2, column=3).style == "highlight"
        assert ws["C2"].style == "highlight"


class TestEquality:
    def test_eq(self):
        a = NamedStyle(name="s", font=Font(bold=True), number_format="0.00")
        b = NamedStyle(name="s", font=Font(bold=True), number_format="0.00")
        c = NamedStyle(name="s", font=Font(bold=False))
        assert a == b
        assert a != c
        assert a != "s"

    def test_repr(self):
        assert "NamedStyle(name='s'" in repr(NamedStyle(name="s"))


class TestSaveReadback:
    def test_named_style_readback(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        wb.add_named_style(
            NamedStyle(
                name="header",
                font=Font(bold=True, size=13, color="FFFFFF"),
                fill=PatternFill(fill_type="solid", start_color="4472C4"),
                alignment=Alignment(horizontal="center", vertical="center"),
                border=Border(bottom=Side(style="medium", color="000000")),
            )
        )
        for col, label in enumerate(["Name", "Qty", "Price"], start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.style = "header"
        ws["A2"] = "Widget"
        path = str(tmp_path / "named.xlsx")
        wb.save(path)

        rb = real_openpyxl.load_workbook(path)
        for col in ("A", "B", "C"):
            c = rb.active[f"{col}1"]
            assert c.font.bold is True
            assert c.font.size == 13
            assert c.fill.patternType == "solid"
            assert c.alignment.horizontal == "center"
            assert c.border.bottom.style == "medium"
        assert rb.active["A2"].font.bold is not True

    def test_number_format_readback(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        wb.add_named_style(NamedStyle(name="pct", number_format="0.00%"))
        ws["A1"] = 0.42
        ws["A1"].style = "pct"
        path = str(tmp_path / "pct.xlsx")
        wb.save(path)

        rb = real_openpyxl.load_workbook(path)
        assert rb.active["A1"].number_format == "0.00%"
        assert rb.active["A1"].value == 0.42
